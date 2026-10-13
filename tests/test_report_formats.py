"""报告三格式（MD/DOCX/XLSX）+ 逐字节可复现（M5 门）。

纪律：三格式共用 `report/model.py` 内容源，故同一底稿在三产物里结论一致；
DOCX/XLSX 经 `docxutil.normalize_ooxml()` 收口为确定性字节（固定 zip 时间戳 +
`create_system=0` + 固定 core.xml 时间戳）。
"""

import zipfile

import pytest

from paudit.engine.pipeline import AuditRequest, compute_audit
from paudit.report import render_to_file
from paudit.report.docxutil import FIXED_ZIP_DATETIME, normalize_ooxml

pytest.importorskip("docx")
pytest.importorskip("openpyxl")

TEMPLATES = ("self", "third_party")


def _draft(demo_ruleset, demo_card):
    return compute_audit(
        AuditRequest(demo_card["input_card"], demo_card["manual_results"],
                     demo_card["audit_scope"]), demo_ruleset).to_dict()


# ---------------------------------------------------------------------------
# 三格式产出
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("ext", ["md", "docx", "xlsx"])
@pytest.mark.parametrize("template", TEMPLATES)
def test_render_to_file_produces_artifact(demo_ruleset, demo_card, tmp_path, ext, template):
    out = tmp_path / ("r.%s" % ext)
    render_to_file(_draft(demo_ruleset, demo_card), demo_ruleset, str(out),
                   template=template)
    assert out.is_file() and out.stat().st_size > 0


def test_render_to_file_rejects_unknown_extension(demo_ruleset, demo_card, tmp_path):
    from paudit.errors import InputError

    with pytest.raises(InputError):
        render_to_file(_draft(demo_ruleset, demo_card), demo_ruleset,
                       str(tmp_path / "r.pdf"))


# ---------------------------------------------------------------------------
# 逐字节可复现（M5 门核心）
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("ext", ["md", "docx", "xlsx"])
@pytest.mark.parametrize("template", TEMPLATES)
def test_report_is_byte_reproducible(demo_ruleset, demo_card, tmp_path, ext, template):
    """同输入两次渲染须逐字节一致（无系统时间戳）。"""
    draft = _draft(demo_ruleset, demo_card)
    a = tmp_path / ("a.%s" % ext)
    b = tmp_path / ("b.%s" % ext)
    render_to_file(draft, demo_ruleset, str(a), template=template)
    render_to_file(draft, demo_ruleset, str(b), template=template)
    assert a.read_bytes() == b.read_bytes(), "报告产物不可复现: .%s/%s" % (ext, template)


def test_docx_zip_entries_have_fixed_timestamp_and_create_system(demo_ruleset, demo_card, tmp_path):
    out = tmp_path / "r.docx"
    render_to_file(_draft(demo_ruleset, demo_card), demo_ruleset, str(out))
    z = zipfile.ZipFile(str(out))
    assert z.infolist(), "docx 包为空"
    for info in z.infolist():
        assert info.date_time == FIXED_ZIP_DATETIME, "条目时间戳未固定: %s" % info.filename
        assert info.create_system == 0, "create_system 未归零: %s" % info.filename


def test_xlsx_zip_entries_have_fixed_timestamp_and_create_system(demo_ruleset, demo_card, tmp_path):
    out = tmp_path / "r.xlsx"
    render_to_file(_draft(demo_ruleset, demo_card), demo_ruleset, str(out))
    z = zipfile.ZipFile(str(out))
    for info in z.infolist():
        assert info.date_time == FIXED_ZIP_DATETIME, "条目时间戳未固定: %s" % info.filename
        assert info.create_system == 0, "create_system 未归零: %s" % info.filename


def test_normalize_ooxml_is_idempotent():
    """收口后再收口不得改变字节（幂等）。"""
    import io

    import docx

    d = docx.Document()
    d.add_paragraph("x")
    buf = io.BytesIO()
    d.save(buf)
    once = normalize_ooxml(buf.getvalue())
    twice = normalize_ooxml(once)
    assert once == twice


# ---------------------------------------------------------------------------
# 内容正确性（不是只有字节稳定，内容也要对）
# ---------------------------------------------------------------------------
def test_docx_contains_sections_and_tables(demo_ruleset, demo_card, tmp_path):
    import docx

    out = tmp_path / "r.docx"
    render_to_file(_draft(demo_ruleset, demo_card), demo_ruleset, str(out),
                   template="third_party")
    d = docx.Document(str(out))
    headings = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    for want in ("一、审计范围", "三、审计触发与频次判定", "四、检查单全项结果",
                 "五、问题清单与整改台账", "六、证据清单与指纹", "七、声明"):
        assert want in headings, "DOCX 缺节: %s" % want
    assert d.tables, "DOCX 无表格"


def test_xlsx_has_fixed_sheets_and_content(demo_ruleset, demo_card, tmp_path):
    import openpyxl

    out = tmp_path / "r.xlsx"
    render_to_file(_draft(demo_ruleset, demo_card), demo_ruleset, str(out))
    wb = openpyxl.load_workbook(str(out))
    assert wb.sheetnames == ["概览", "审计依据", "触发判定", "检查单", "问题台账", "证据清单"]
    assert wb["概览"]["A2"].value == "报告标题"
    # 检查单表须含每个检查项 id
    checklist_text = "\n".join(
        str(c.value) for row in wb["检查单"].iter_rows() for c in row if c.value)
    for item in demo_ruleset.checklist:
        assert item.id in checklist_text, "XLSX 检查单缺项: %s" % item.id


def test_three_formats_agree_on_conclusions(demo_ruleset, demo_card, tmp_path):
    """三格式结论一致：检查单四态须同时出现在 MD/DOCX/XLSX 产物里。"""
    import docx
    import openpyxl

    draft = _draft(demo_ruleset, demo_card)
    md_path = tmp_path / "r.md"
    docx_path = tmp_path / "r.docx"
    xlsx_path = tmp_path / "r.xlsx"
    render_to_file(draft, demo_ruleset, str(md_path))
    render_to_file(draft, demo_ruleset, str(docx_path))
    render_to_file(draft, demo_ruleset, str(xlsx_path))

    md_text = md_path.read_text(encoding="utf-8")
    docx_text = "\n".join(
        [p.text for p in docx.Document(str(docx_path)).paragraphs]
        + [c.text for t in docx.Document(str(docx_path)).tables
           for row in t.rows for c in row.cells])
    wb = openpyxl.load_workbook(str(xlsx_path))
    xlsx_text = "\n".join(
        str(c.value) for row in wb["检查单"].iter_rows() for c in row if c.value)

    for it in draft["item_results"]:
        assert it["item_id"] in md_text
        assert it["item_id"] in docx_text
        assert it["item_id"] in xlsx_text
        assert it["state"] in md_text
        assert it["state"] in docx_text
        assert it["state"] in xlsx_text
