"""DOCX 报告渲染（optional extra `[report]`；python-docx 只在函数体内导入）。

内容来自 `report/model.py`（单一内容源），本模块只做"块 → Word 语法"的映射，
不含任何判定逻辑。落盘经 `docxutil.write_ooxml()` 收口为确定性字节。
"""

from .model import Bullets, Heading, Note, Para, Table, build


def render_docx(draft, ruleset, out_path, template="self"):
    """渲染 DOCX 报告并落盘（确定性字节）。返回输出路径。"""
    import docx  # 函数体内导入：核心层 AST 守门（G3）要求 report 包不得顶层依赖重件
    from docx.shared import Pt

    from .docxutil import write_ooxml

    report = build(draft, ruleset, template=template)
    doc = docx.Document()

    doc.add_heading(report.title, level=0)
    coord = doc.add_paragraph("规则包坐标：%s" % report.coord)
    coord.runs[0].bold = True

    for section in report.sections:
        doc.add_heading(section.title, level=1)
        for block in section.blocks:
            _add_block(doc, block)

    for line in report.signature_lines:
        doc.add_paragraph(line)

    buf = _save_bytes(doc)
    return write_ooxml(buf, out_path)


def _add_block(doc, block):
    if isinstance(block, Heading):
        doc.add_heading(block.text, level=min(block.level, 4))
    elif isinstance(block, Para):
        doc.add_paragraph(block.text)
    elif isinstance(block, Bullets):
        for item in block.items:
            doc.add_paragraph(item, style="List Bullet")
    elif isinstance(block, Note):
        p = doc.add_paragraph(block.text)
        p.runs[0].italic = True if p.runs else None
    elif isinstance(block, Table):
        _add_table(doc, block)
    else:  # 未知块：不静默丢弃，落为段落（保内容不丢）
        doc.add_paragraph(str(getattr(block, "text", block)))


def _add_table(doc, table):
    if not table.headers:
        return
    t = doc.add_table(rows=1, cols=len(table.headers))
    t.style = "Table Grid"
    for i, h in enumerate(table.headers):
        cell = t.rows[0].cells[i]
        cell.text = str(h)
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for row in table.rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            if i < len(cells):
                cells[i].text = "" if v is None else str(v)
    if not table.rows:  # 空表补一行"（无）"，避免 Word 里只留表头
        cells = t.add_row().cells
        cells[0].text = "（无）"


def _save_bytes(doc):
    """把 Document 存为 bytes（python-docx 需要流；统一走 BytesIO）。"""
    import io

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
